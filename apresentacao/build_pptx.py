"""
Doutor AI · Apresentação Institucional · PPTX builder
Reproduz, de forma editável, os 12 slides do PDF aplicando a paleta,
tipografia e estrutura do Branding Book v1.0.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# =============================================================
# DESIGN TOKENS  (extraídos do Branding Book v1.0)
# =============================================================
C_BLUE_DEEP  = RGBColor(0x00, 0x3D, 0x5B)
C_BLUE_BRAND = RGBColor(0x00, 0x71, 0xBC)
C_BLUE_MID   = RGBColor(0x3B, 0x8F, 0xE0)
C_BLUE_LIGHT = RGBColor(0x8F, 0xC3, 0xEA)
C_BLUE_PALE  = RGBColor(0xEB, 0xF1, 0xF8)

C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_PAPER   = RGBColor(0xFA, 0xFA, 0xFA)
C_GRAY100 = RGBColor(0xF1, 0xF2, 0xF4)
C_GRAY300 = RGBColor(0xD1, 0xD5, 0xDB)
C_GRAY500 = RGBColor(0x9C, 0xA3, 0xAF)
C_GRAY700 = RGBColor(0x4B, 0x55, 0x63)
C_CHUMBO  = RGBColor(0x1A, 0x1F, 0x2E)

FONT_SANS = "Inter"
FONT_MONO = "JetBrains Mono"

# =============================================================
# PRESENTATION SETUP · 16:9
# =============================================================
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW = prs.slide_width
SH = prs.slide_height

blank_layout = prs.slide_layouts[6]   # truly blank

# Standard margins (mirror the HTML deck)
M_LEFT  = Inches(0.7)
M_RIGHT = Inches(0.7)
M_TOP   = Inches(0.55)


# =============================================================
# HELPERS
# =============================================================
def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=None):
    """Add a filled/outlined rectangle. fill/line can be None for no fill/line."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        if line_w is not None:
            shape.line.width = line_w
    return shape


def add_line(slide, x, y, w, color, thickness=Emu(7000)):
    """Horizontal line (thin rectangle)."""
    return add_rect(slide, x, y, w, thickness, fill=color, line=None)


def add_text(slide, x, y, w, h, runs, *,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             margin=(Pt(0), Pt(0), Pt(0), Pt(0)),
             line_spacing=None, fill=None):
    """
    runs: list of dicts: {text, font, size, bold, italic, color, spacing}
    """
    tb = slide.shapes.add_textbox(x, y, w, h)
    if fill is not None:
        tb.fill.solid(); tb.fill.fore_color.rgb = fill
    tb.line.fill.background()
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = margin
    tf.vertical_anchor = anchor

    # First paragraph already exists
    first = True
    for r in runs:
        if r.get("new_para") or first:
            if first:
                p = tf.paragraphs[0]
                first = False
            else:
                p = tf.add_paragraph()
            p.alignment = r.get("align", align)
            if line_spacing is not None:
                p.line_spacing = line_spacing
        if r.get("text", "") == "" and r.get("new_para"):
            # blank para spacer
            continue
        run = p.add_run()
        run.text = r["text"]
        f = run.font
        f.name = r.get("font", FONT_SANS)
        f.size = r.get("size", Pt(11))
        f.bold = r.get("bold", False)
        f.italic = r.get("italic", False)
        col = r.get("color", C_CHUMBO)
        f.color.rgb = col
        if r.get("spacing") is not None:
            # letter spacing via XML (spc in 1/100 pt)
            rPr = run._r.get_or_add_rPr()
            rPr.set("spc", str(int(r["spacing"] * 100)))
    return tb


def kicker(slide, x, y, text, *, color=C_BLUE_BRAND, w=Inches(8)):
    add_text(slide, x, y, w, Inches(0.22),
        runs=[{
            "text": "▸  " + text.upper(),
            "font": FONT_MONO,
            "size": Pt(9),
            "color": color,
            "spacing": 1.6,
            "bold": False,
        }],
        anchor=MSO_ANCHOR.TOP,
        margin=(Pt(0), Pt(0), Pt(0), Pt(0)),
    )


def header(slide, section_label, *, dark=False):
    """Reusable header: logo + Doutor-AI · Apresentação Institucional | section_label"""
    fg = C_WHITE if dark else C_CHUMBO
    muted = RGBColor(0xFF, 0xFF, 0xFF) if dark else C_GRAY500
    sep = RGBColor(0xFF, 0xFF, 0xFF) if dark else C_GRAY500
    # logo D mark as filled circle stand-in (simplified) -- to keep PPTX clean
    # We draw a small square with the "D" character as bold text inside a square
    lx, ly = M_LEFT, M_TOP - Inches(0.05)
    lsize = Inches(0.22)
    # Logo: small rounded square with "D"
    logo_color = C_WHITE if dark else C_BLUE_BRAND
    logo = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, lx, ly, lsize, lsize)
    logo.shadow.inherit = False
    logo.fill.solid(); logo.fill.fore_color.rgb = logo_color
    logo.line.fill.background()
    # Adjust rounding (smaller corner)
    logo.adjustments[0] = 0.18
    # "D" letter on top
    add_text(slide, lx, ly - Inches(0.005), lsize, lsize,
        runs=[{"text": "D", "font": FONT_SANS, "size": Pt(11), "bold": True,
               "color": C_WHITE if not dark else C_BLUE_BRAND, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        margin=(Pt(0), Pt(0), Pt(0), Pt(0)))

    # Wordmark + section
    tx = lx + lsize + Inches(0.08)
    add_text(slide, tx, M_TOP - Inches(0.03),
             Inches(7), Inches(0.28),
        runs=[
            {"text": "Doutor-AI", "font": FONT_SANS, "size": Pt(9),
             "bold": True, "color": fg},
            {"text": "  ·  ", "font": FONT_SANS, "size": Pt(9), "color": muted},
            {"text": "Apresentação Institucional", "font": FONT_MONO, "size": Pt(7.5),
             "color": muted, "spacing": 1.4},
        ],
        anchor=MSO_ANCHOR.MIDDLE,
        margin=(Pt(0), Pt(0), Pt(0), Pt(0)),
    )

    # Section label on the right
    add_text(slide, SW - M_RIGHT - Inches(4),
             M_TOP - Inches(0.03), Inches(4), Inches(0.28),
        runs=[{"text": section_label.upper(), "font": FONT_MONO, "size": Pt(7.5),
               "color": muted, "spacing": 1.4, "align": PP_ALIGN.RIGHT}],
        align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
        margin=(Pt(0), Pt(0), Pt(0), Pt(0)),
    )

    # Thin horizontal rule under header
    rule_y = M_TOP + Inches(0.27)
    rule_color = RGBColor(0xFF,0xFF,0xFF) if dark else C_GRAY300
    add_line(slide, M_LEFT, rule_y, SW - M_LEFT - M_RIGHT, rule_color, thickness=Emu(7000))


def footer(slide, page_num, *, dark=False):
    fy = SH - Inches(0.4)
    fg = RGBColor(0xFF,0xFF,0xFF) if dark else C_GRAY700
    muted = RGBColor(0xFF,0xFF,0xFF) if dark else C_GRAY500
    add_text(slide, M_LEFT, fy, Inches(6), Inches(0.2),
        runs=[{"text": "DOUTOR-AI · V1.0 · 2026", "font": FONT_MONO, "size": Pt(7.5),
               "color": muted, "spacing": 1.4}],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_text(slide, SW - M_RIGHT - Inches(2), fy, Inches(2), Inches(0.2),
        runs=[{"text": f"{page_num:02d} / 12", "font": FONT_MONO, "size": Pt(7.5),
               "color": fg, "spacing": 1.4, "align": PP_ALIGN.RIGHT}],
        align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
    )


def headline(slide, x, y, w, parts, *, size=Pt(36), line_spacing=1.02):
    """parts: list of (text, italic_blue_bool). Joined as one paragraph."""
    runs = []
    for i, (t, italic) in enumerate(parts):
        run = {
            "text": t,
            "font": FONT_SANS,
            "size": size,
            "bold": True,
            "color": C_BLUE_BRAND if italic else C_CHUMBO,
            "italic": italic,
            "spacing": -0.6,
        }
        runs.append(run)
    add_text(slide, x, y, w, Inches(2), runs=runs, line_spacing=line_spacing,
             margin=(Pt(0), Pt(0), Pt(0), Pt(0)))


def body_text(slide, x, y, w, h, text, *, size=Pt(12), color=C_GRAY700, line_spacing=1.45,
              bold_terms=None):
    """Body paragraph with optional bold highlights."""
    bold_terms = bold_terms or []
    runs = []
    rest = text
    if bold_terms:
        # Build runs splitting on bold terms
        import re
        pattern = "|".join(re.escape(t) for t in bold_terms)
        last = 0
        for m in re.finditer(pattern, rest):
            if m.start() > last:
                runs.append({"text": rest[last:m.start()], "font": FONT_SANS,
                             "size": size, "color": color})
            runs.append({"text": m.group(), "font": FONT_SANS, "size": size,
                         "bold": True, "color": C_CHUMBO})
            last = m.end()
        if last < len(rest):
            runs.append({"text": rest[last:], "font": FONT_SANS, "size": size, "color": color})
    else:
        runs.append({"text": rest, "font": FONT_SANS, "size": size, "color": color})
    add_text(slide, x, y, w, h, runs=runs, line_spacing=line_spacing,
             margin=(Pt(0), Pt(0), Pt(0), Pt(0)))


def card(slide, x, y, w, h, *, fill=C_WHITE, border=C_GRAY300, dashed=False,
         left_accent=None):
    rect = add_rect(slide, x, y, w, h, fill=fill, line=border, line_w=Emu(6000))
    if dashed:
        ln = rect.line._get_or_add_ln()
        # Set dash style
        for pr in ln.findall(qn('a:prstDash')):
            ln.remove(pr)
        dash = etree.SubElement(ln, qn('a:prstDash'))
        dash.set('val', 'dash')
    if left_accent is not None:
        add_rect(slide, x, y, Emu(28000), h, fill=left_accent, line=None)
    return rect


def card_num_title_dsc(slide, x, y, w, h, num_label, title, dsc, *,
                       title_size=Pt(14), dsc_size=Pt(9.5),
                       num_color=C_BLUE_BRAND):
    card(slide, x, y, w, h)
    pad = Inches(0.18)
    # num label
    add_text(slide, x + pad, y + pad, w - 2*pad, Inches(0.22),
        runs=[{"text": num_label.upper(), "font": FONT_MONO, "size": Pt(8),
               "color": num_color, "spacing": 1.6, "bold": False}])
    add_text(slide, x + pad, y + pad + Inches(0.28), w - 2*pad, Inches(0.42),
        runs=[{"text": title, "font": FONT_SANS, "size": title_size,
               "bold": True, "color": C_CHUMBO, "spacing": -0.4}],
        line_spacing=1.15)
    add_text(slide, x + pad, y + pad + Inches(0.72), w - 2*pad, h - Inches(0.95),
        runs=[{"text": dsc, "font": FONT_SANS, "size": dsc_size,
               "color": C_GRAY700}],
        line_spacing=1.45)


# =============================================================
# SLIDE 1 · CAPA
# =============================================================
s = prs.slides.add_slide(blank_layout)
# Solid brand-blue background
bg = add_rect(s, 0, 0, SW, SH, fill=C_BLUE_BRAND, line=None)

header(s, "v1.0 · 2026", dark=True)
# Override section right-side text (already added by header). Add a strong custom one
# by drawing on top — we keep header's section as v1.0 · 2026.

# Kicker
kicker(s, M_LEFT, Inches(1.85), "APRESENTAÇÃO INSTITUCIONAL · INTERNA", color=C_WHITE)

# Hero
add_text(s, M_LEFT, Inches(2.15), SW - M_LEFT - M_RIGHT, Inches(1.8),
    runs=[{"text": "Doutor AI.", "font": FONT_SANS, "size": Pt(96),
           "bold": True, "color": C_WHITE, "spacing": -4.5}],
    line_spacing=0.98)

# Rule
add_line(s, M_LEFT, Inches(4.2), Inches(3.5),
         RGBColor(0xFF,0xFF,0xFF), thickness=Emu(7000))

# Subtitle (mixed runs: italic accent for "jornada")
add_text(s, M_LEFT, Inches(4.4), SW - M_LEFT - M_RIGHT, Inches(0.9),
    runs=[
        {"text": "Inteligência artificial para transformar\n",
         "font": FONT_SANS, "size": Pt(28), "bold": False,
         "color": C_WHITE, "spacing": -0.6},
        {"text": "a ", "font": FONT_SANS, "size": Pt(28), "bold": False,
         "color": C_WHITE, "spacing": -0.6, "new_para": True},
        {"text": "jornada", "font": FONT_SANS, "size": Pt(28), "bold": True,
         "italic": True, "color": C_BLUE_LIGHT, "spacing": -0.6},
        {"text": " de saúde.", "font": FONT_SANS, "size": Pt(28),
         "color": C_WHITE, "spacing": -0.6},
    ], line_spacing=1.1)

# Lead paragraph
add_text(s, M_LEFT, Inches(5.85), Inches(8.5), Inches(0.8),
    runs=[{"text": "Mais eficiência, inteligência operacional e suporte clínico para hospitais, operadoras e clínicas.",
           "font": FONT_SANS, "size": Pt(13),
           "color": RGBColor(0xFF,0xFF,0xFF)}],
    line_spacing=1.45)

# Meta footer (Documento / Uso / Slides)
add_line(s, M_LEFT, SH - Inches(0.95), SW - M_LEFT - M_RIGHT,
         RGBColor(0xFF,0xFF,0xFF), thickness=Emu(7000))
meta_y = SH - Inches(0.82)
labels = [("DOCUMENTO", "Institucional"),
          ("USO",       "Interno · Comercial"),
          ("SLIDES",    "12")]
mx = M_LEFT
for lab, val in labels:
    add_text(s, mx, meta_y, Inches(2.2), Inches(0.2),
        runs=[{"text": lab, "font": FONT_MONO, "size": Pt(7.5),
               "color": RGBColor(0xFF,0xFF,0xFF), "spacing": 1.6}])
    add_text(s, mx, meta_y + Inches(0.22), Inches(2.2), Inches(0.25),
        runs=[{"text": val, "font": FONT_SANS, "size": Pt(10.5),
               "bold": True, "color": C_WHITE}])
    mx += Inches(2.4)


# =============================================================
# SLIDE 2 · QUEM SOMOS
# =============================================================
s = prs.slides.add_slide(blank_layout)
header(s, "01 · Quem somos")
kicker(s, M_LEFT, Inches(1.05), "QUEM SOMOS")

# Headline
add_text(s, M_LEFT, Inches(1.35), Inches(12), Inches(1.7),
    runs=[
        {"text": "A camada de ", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "color": C_CHUMBO, "spacing": -1},
        {"text": "inteligência", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1},
        {"text": "​", "font": FONT_SANS, "size": Pt(36)},
        {"text": "assistencial e operacional.", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "color": C_CHUMBO, "spacing": -1, "new_para": True},
    ],
    line_spacing=1.0,
)

body_text(s, M_LEFT, Inches(3.0), Inches(11.8), Inches(1.0),
    "A Doutor AI é uma health tech que aplica inteligência artificial em jornadas "
    "assistenciais e operacionais de saúde — conectando pacientes, profissionais e "
    "instituições em fluxos mais eficientes, seguros e inteligentes.",
    size=Pt(13), bold_terms=["Doutor AI", "eficientes, seguros e inteligentes"])

# Pilares
kicker(s, M_LEFT, Inches(4.35), "TRÊS PILARES")
pillars = [
    ("01 · JORNADA",  "IA aplicada à jornada de saúde",
     "Conversas estruturadas, coleta clínica e orquestração de fluxos do pré ao pós-atendimento."),
    ("02 · AUTOMAÇÃO","Automação de tarefas operacionais",
     "Redução de retrabalho em processos repetitivos, registros e integrações entre sistemas."),
    ("03 · SUPORTE",  "Suporte à decisão e ao atendimento",
     "Copiloto clínico, alertas e protocolos para apoiar — não substituir — o profissional de saúde."),
]
cw = (SW - M_LEFT - M_RIGHT - Inches(0.3)) / 3
cy = Inches(4.7)
cx = M_LEFT
for num, ttl, dsc in pillars:
    card_num_title_dsc(s, cx, cy, cw, Inches(1.95), num, ttl, dsc)
    cx += cw + Inches(0.15)

footer(s, 2)


# =============================================================
# SLIDE 3 · O PROBLEMA (6 cards)
# =============================================================
s = prs.slides.add_slide(blank_layout)
header(s, "02 · Problema")
kicker(s, M_LEFT, Inches(1.05), "O PROBLEMA QUE RESOLVEMOS")

add_text(s, M_LEFT, Inches(1.35), Inches(12), Inches(1.7),
    runs=[
        {"text": "Saúde opera no limite.", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "color": C_CHUMBO, "spacing": -1},
        {"text": "Todos os dias.", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1,
         "new_para": True},
    ], line_spacing=1.0)

pains = [
    ("01 · DEMANDA",   "Volume e filas",       "Alto volume de atendimentos e filas que excedem a capacidade instalada."),
    ("02 · PROCESSOS", "Manual e repetitivo",  "Processos manuais, repetitivos e sem padronização entre unidades e equipes."),
    ("03 · DADOS",     "Perdas entre etapas",  "Informações clínicas que se perdem ao longo da jornada e prejudicam decisões."),
    ("04 · EQUIPE",    "Sobrecarga clínica",   "Profissionais sobrecarregados em tarefas administrativas, com menos foco no paciente."),
    ("05 · JORNADA",   "Acompanhamento frágil","Dificuldade de acompanhar pacientes de forma ativa entre consultas e exames."),
    ("06 · NEGÓCIO",   "Custo & receita",      "Ineficiências que impactam custo, receita, experiência e qualidade assistencial."),
]
cw = (SW - M_LEFT - M_RIGHT - Inches(0.3)) / 3
ch = Inches(1.55)
cy0 = Inches(3.95)
for i, (n, t, d) in enumerate(pains):
    row, col = divmod(i, 3)
    cx = M_LEFT + col * (cw + Inches(0.15))
    cy = cy0 + row * (ch + Inches(0.18))
    card_num_title_dsc(s, cx, cy, cw, ch, n, t, d,
                       title_size=Pt(13), dsc_size=Pt(9))

footer(s, 3)


# =============================================================
# SLIDE 4 · PROPOSTA DE VALOR · ANTES / DEPOIS
# =============================================================
s = prs.slides.add_slide(blank_layout)
header(s, "03 · Proposta de valor")
kicker(s, M_LEFT, Inches(1.05), "PROPOSTA DE VALOR")

add_text(s, M_LEFT, Inches(1.35), Inches(12), Inches(1.7),
    runs=[
        {"text": "IA para dar ", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "color": C_CHUMBO, "spacing": -1},
        {"text": "escala", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1},
        {"text": "sem perder qualidade.", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "color": C_CHUMBO, "spacing": -1, "new_para": True},
    ], line_spacing=1.0)

body_text(s, M_LEFT, Inches(3.0), Inches(11.8), Inches(0.9),
    "A Doutor AI estrutura conversas, coleta dados, organiza informações clínicas, "
    "automatiza etapas operacionais e apoia profissionais em tempo real.",
    size=Pt(13), bold_terms=["em tempo real"])

# Before / After columns
col_w = (SW - M_LEFT - M_RIGHT - Inches(0.3)) / 2
col_h = Inches(2.95)
col_y = Inches(4.0)

# BEFORE
card(s, M_LEFT, col_y, col_w, col_h, fill=C_PAPER, border=C_GRAY300)
add_text(s, M_LEFT + Inches(0.25), col_y + Inches(0.25), col_w, Inches(0.25),
    runs=[{"text": "▸  ANTES", "font": FONT_MONO, "size": Pt(9),
           "color": C_GRAY700, "spacing": 1.8, "bold": True}])
before_items = [
    "Atendimento fragmentado entre canais",
    "Coleta clínica manual e inconsistente",
    "Baixa previsibilidade da jornada",
    "Retrabalho entre etapas e sistemas",
    "Decisões sem suporte estruturado",
]
iy = col_y + Inches(0.6)
for i, it in enumerate(before_items):
    add_text(s, M_LEFT + Inches(0.25), iy, col_w - Inches(0.5), Inches(0.32),
        runs=[{"text": it, "font": FONT_SANS, "size": Pt(11.5),
               "color": C_GRAY700}])
    if i < len(before_items) - 1:
        add_line(s, M_LEFT + Inches(0.25), iy + Inches(0.4),
                 col_w - Inches(0.5), C_GRAY300, thickness=Emu(6000))
    iy += Inches(0.42)

# AFTER
ax = M_LEFT + col_w + Inches(0.3)
card(s, ax, col_y, col_w, col_h, fill=C_BLUE_PALE, border=C_BLUE_LIGHT)
add_text(s, ax + Inches(0.25), col_y + Inches(0.25), col_w, Inches(0.25),
    runs=[{"text": "▸  DEPOIS", "font": FONT_MONO, "size": Pt(9),
           "color": C_BLUE_BRAND, "spacing": 1.8, "bold": True}])
after_items = [
    ("Jornada estruturada", " e orquestrada por IA"),
    ("Dados organizados",   " e padronizados em tempo real"),
    ("Protocolos assistidos", " e alertas inteligentes"),
    ("Equipe mais produtiva", " e focada no paciente"),
    ("Visibilidade ponta-a-ponta", " da operação"),
]
iy = col_y + Inches(0.6)
for i, (b, r) in enumerate(after_items):
    add_text(s, ax + Inches(0.25), iy, col_w - Inches(0.5), Inches(0.32),
        runs=[
            {"text": b, "font": FONT_SANS, "size": Pt(11.5),
             "bold": True, "color": C_CHUMBO},
            {"text": r, "font": FONT_SANS, "size": Pt(11.5),
             "color": C_CHUMBO},
        ])
    if i < len(after_items) - 1:
        add_line(s, ax + Inches(0.25), iy + Inches(0.4),
                 col_w - Inches(0.5), C_BLUE_LIGHT, thickness=Emu(6000))
    iy += Inches(0.42)

footer(s, 4)


# =============================================================
# SLIDE 5 · CAPACIDADES (8 cards 4×2)
# =============================================================
s = prs.slides.add_slide(blank_layout)
header(s, "04 · Capacidades")
kicker(s, M_LEFT, Inches(1.05), "O QUE A PLATAFORMA FAZ")
add_text(s, M_LEFT, Inches(1.35), Inches(12.5), Inches(1.3),
    runs=[
        {"text": "Capacidades aplicadas a ", "font": FONT_SANS, "size": Pt(34),
         "bold": True, "color": C_CHUMBO, "spacing": -1},
        {"text": "fluxos clínicos reais.", "font": FONT_SANS, "size": Pt(34),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1},
    ], line_spacing=1.0)

caps = [
    ("Triagem inteligente",      "Por texto, voz ou canais digitais — com priorização clínica."),
    ("Coleta clínica estruturada","Dados organizados desde o primeiro contato com o paciente."),
    ("Classificação de risco",   "Direcionamento assistencial baseado em protocolos."),
    ("Transcrição & sumarização","Registro automático de atendimentos com resumo clínico."),
    ("Copiloto clínico",         "Suporte em tempo real ao raciocínio e à conduta."),
    ("Integração com prontuário","Conexão com EHR e sistemas hospitalares existentes."),
    ("Alertas & protocolos",     "Apoio baseado em diretrizes institucionais e clínicas."),
    ("Navegação de pacientes",   "Acompanhamento ativo da jornada entre etapas."),
]
cw = (SW - M_LEFT - M_RIGHT - Inches(0.4)) / 4
ch = Inches(1.25)
cy0 = Inches(3.0)
for i, (t, d) in enumerate(caps):
    row, col = divmod(i, 4)
    cx = M_LEFT + col * (cw + Inches(0.13))
    cy = cy0 + row * (ch + Inches(0.15))
    card(s, cx, cy, cw, ch)
    pad = Inches(0.15)
    add_text(s, cx + pad, cy + pad, cw - 2*pad, Inches(0.38),
        runs=[{"text": t, "font": FONT_SANS, "size": Pt(12),
               "bold": True, "color": C_CHUMBO, "spacing": -0.3}],
        line_spacing=1.15)
    add_text(s, cx + pad, cy + pad + Inches(0.45), cw - 2*pad, ch - Inches(0.65),
        runs=[{"text": d, "font": FONT_SANS, "size": Pt(8.5),
               "color": C_GRAY700}], line_spacing=1.4)

# Callout
cy = Inches(5.9)
ch = Inches(0.78)
card(s, M_LEFT, cy, SW - M_LEFT - M_RIGHT, ch,
     fill=C_BLUE_PALE, border=C_BLUE_PALE, left_accent=C_BLUE_BRAND)
add_text(s, M_LEFT + Inches(0.25), cy + Inches(0.12), Inches(3), Inches(0.22),
    runs=[{"text": "▸  MULTIMODAL", "font": FONT_MONO, "size": Pt(8.5),
           "color": C_BLUE_BRAND, "spacing": 1.8, "bold": True}])
add_text(s, M_LEFT + Inches(0.25), cy + Inches(0.36),
         SW - M_LEFT - M_RIGHT - Inches(0.5), Inches(0.4),
    runs=[
        {"text": "Capacidades operam de forma combinada em ", "font": FONT_SANS,
         "size": Pt(10.5), "color": C_CHUMBO},
        {"text": "texto, voz e dados estruturados", "font": FONT_SANS,
         "size": Pt(10.5), "color": C_CHUMBO, "bold": True},
        {"text": ", sempre com validação humana nos pontos críticos da jornada.",
         "font": FONT_SANS, "size": Pt(10.5), "color": C_CHUMBO},
    ], line_spacing=1.4)

footer(s, 5)


# =============================================================
# SLIDE 6 · JORNADAS (8 cards 4×2)
# =============================================================
s = prs.slides.add_slide(blank_layout)
header(s, "05 · Jornadas")
kicker(s, M_LEFT, Inches(1.05), "JORNADAS ATENDIDAS")
add_text(s, M_LEFT, Inches(1.35), Inches(12), Inches(1.3),
    runs=[
        {"text": "Onde a Doutor AI ", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "color": C_CHUMBO, "spacing": -1},
        {"text": "atua.", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1},
    ], line_spacing=1.0)

journeys = [
    ("01 · PA / PS",       "Pronto atendimento e pronto-socorro digital"),
    ("02 · TELE",          "Telemedicina e atendimento à distância"),
    ("03 · PRÉ-CONSULTA",  "Triagem e coleta antes da consulta"),
    ("04 · PÓS-CONSULTA",  "Checkout médico e orientação pós-atendimento"),
    ("05 · CRÔNICOS",      "Gestão e acompanhamento de pacientes crônicos"),
    ("06 · FATURAMENTO",   "Apoio a glosas e pertinência clínica"),
    ("07 · REGISTROS",     "Automação de documentação clínica"),
    ("08 · CANAIS",        "Atendimento via canais digitais e mensageria"),
]
cw = (SW - M_LEFT - M_RIGHT - Inches(0.4)) / 4
ch = Inches(1.15)
cy0 = Inches(3.1)
for i, (n, t) in enumerate(journeys):
    row, col = divmod(i, 4)
    cx = M_LEFT + col * (cw + Inches(0.13))
    cy = cy0 + row * (ch + Inches(0.18))
    card(s, cx, cy, cw, ch)
    pad = Inches(0.16)
    add_text(s, cx + pad, cy + pad, cw - 2*pad, Inches(0.22),
        runs=[{"text": n, "font": FONT_MONO, "size": Pt(8),
               "color": C_BLUE_BRAND, "spacing": 1.6, "bold": True}])
    add_text(s, cx + pad, cy + pad + Inches(0.3), cw - 2*pad, ch - Inches(0.55),
        runs=[{"text": t, "font": FONT_SANS, "size": Pt(11),
               "bold": True, "color": C_CHUMBO, "spacing": -0.3}],
        line_spacing=1.2)

body_text(s, M_LEFT, Inches(5.9), SW - M_LEFT - M_RIGHT, Inches(0.6),
    "Cada jornada pode ser implementada de forma independente ou combinada, "
    "iniciando por uma frente de alto impacto e expandindo conforme a maturidade da operação.",
    size=Pt(10.5), bold_terms=["iniciando por uma frente de alto impacto"])

footer(s, 6)


# =============================================================
# SLIDE 7 · CLIENTES E PERFIS
# =============================================================
s = prs.slides.add_slide(blank_layout)
header(s, "06 · Clientes & parceiros")
kicker(s, M_LEFT, Inches(1.05), "CLIENTES & PARCEIROS")
add_text(s, M_LEFT, Inches(1.35), Inches(12.5), Inches(1.7),
    runs=[
        {"text": "Construído com ", "font": FONT_SANS, "size": Pt(34),
         "bold": True, "color": C_CHUMBO, "spacing": -1},
        {"text": "instituições", "font": FONT_SANS, "size": Pt(34),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1},
        {"text": "de referência.", "font": FONT_SANS, "size": Pt(34),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1,
         "new_para": True},
    ], line_spacing=1.0)

body_text(s, M_LEFT, Inches(3.05), Inches(11.8), Inches(0.7),
    "Aplicações em hospitais, redes, operadoras e telemedicina "
    "— em fluxos assistenciais e operacionais.",
    size=Pt(12.5),
    bold_terms=["hospitais, redes, operadoras e telemedicina"])

# 4 client cards
clients = [
    ("REDE HOSPITALAR",         "Rede D'Or São Luiz",
     "Jornadas assistenciais e operacionais em ambiente hospitalar."),
    ("REDE CREDENCIADA",        "Evidia",
     "Do atendimento ao credenciado — auditoria e faturamento."),
    ("TELEMEDICINA",            "Topmed",
     "Telemedicina e atendimento à distância em escala."),
    ("OPERADORA VERTICALIZADA", "Leve Saúde",
     "Operadora verticalizada com modelo próprio de cuidado."),
]
cw = (SW - M_LEFT - M_RIGHT - Inches(0.36)) / 4
ch = Inches(1.85)
cy = Inches(3.95)
for i, (seg, name, dsc) in enumerate(clients):
    cx = M_LEFT + i * (cw + Inches(0.12))
    card(s, cx, cy, cw, ch, fill=C_WHITE, border=C_GRAY300)
    pad = Inches(0.18)
    add_text(s, cx + pad, cy + pad, cw - 2*pad, Inches(0.24),
        runs=[{"text": seg, "font": FONT_MONO, "size": Pt(8),
               "color": C_BLUE_BRAND, "spacing": 1.7, "bold": True}])
    add_text(s, cx + pad, cy + pad + Inches(0.3), cw - 2*pad, Inches(0.6),
        runs=[{"text": name, "font": FONT_SANS, "size": Pt(14),
               "bold": True, "color": C_CHUMBO, "spacing": -0.5}],
        line_spacing=1.1)
    add_text(s, cx + pad, cy + pad + Inches(0.95), cw - 2*pad, ch - Inches(1.15),
        runs=[{"text": dsc, "font": FONT_SANS, "size": Pt(9),
               "color": C_GRAY700}],
        line_spacing=1.4)

# Scale insight card
sy = Inches(5.95)
sh = Inches(1.05)
sx = M_LEFT
sw = SW - M_LEFT - M_RIGHT
card(s, sx, sy, sw, sh, fill=C_WHITE, border=C_GRAY300)
pad = Inches(0.28)
# Big +25
add_text(s, sx + pad, sy + Inches(0.05), Inches(1.6), sh,
    runs=[{"text": "+25", "font": FONT_SANS, "size": Pt(38),
           "bold": True, "color": C_CHUMBO, "spacing": -2}],
    anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
# Title block
tx = sx + Inches(1.95)
tw = sw - Inches(2.2)
add_text(s, tx, sy + Inches(0.16), tw, Inches(0.3),
    runs=[{"text": "unidades / operadoras", "font": FONT_SANS, "size": Pt(15),
           "bold": True, "color": C_CHUMBO, "spacing": -0.4}],
    line_spacing=1.0)
add_text(s, tx, sy + Inches(0.42), tw, Inches(0.28),
    runs=[{"text": "Operadoras verticalizadas e redes regionais",
           "font": FONT_SANS, "size": Pt(10.5), "color": C_GRAY700}])
add_text(s, tx, sy + Inches(0.68), tw, Inches(0.28),
    runs=[{"text": "Replicação do modelo de prontuário-com-IA em rede credenciada",
           "font": FONT_SANS, "size": Pt(11),
           "bold": True, "color": C_BLUE_BRAND, "spacing": -0.3}])

footer(s, 7)


# =============================================================
# SLIDE 8 · DORES POR PÚBLICO INTERNO (4 colunas)
# =============================================================
s = prs.slides.add_slide(blank_layout)
header(s, "07 · Valor por área")
kicker(s, M_LEFT, Inches(1.05), "VALOR POR PÚBLICO INTERNO")
add_text(s, M_LEFT, Inches(1.35), Inches(12), Inches(1.7),
    runs=[
        {"text": "Uma plataforma.", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "color": C_CHUMBO, "spacing": -1},
        {"text": "Quatro públicos.", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1,
         "new_para": True},
    ], line_spacing=1.0)

audiences = [
    ("Operações",  ["Redução de retrabalho", "Padronização do fluxo", "Ganho de produtividade"]),
    ("Assistência",["Melhor coleta de dados clínicos", "Apoio à decisão em tempo real", "Alertas e protocolos institucionais"]),
    ("Gestão",     ["Mais visibilidade da operação", "Indicadores da jornada", "Eficiência operacional"]),
    ("Financeiro", ["Redução de perdas", "Melhor documentação clínica", "Apoio a glosas e pertinência"]),
]
cw = (SW - M_LEFT - M_RIGHT - Inches(0.45)) / 4
ch = Inches(2.7)
cy = Inches(4.0)
for i, (title, items) in enumerate(audiences):
    cx = M_LEFT + i * (cw + Inches(0.15))
    card(s, cx, cy, cw, ch, fill=C_WHITE, border=C_GRAY300)
    pad = Inches(0.2)
    # Dot + title (head)
    dot = slide_shape = None
    add_rect(s, cx + pad, cy + pad + Inches(0.07),
             Inches(0.08), Inches(0.08), fill=C_BLUE_BRAND, line=None)
    add_text(s, cx + pad + Inches(0.16), cy + pad - Inches(0.02),
             cw - 2*pad - Inches(0.16), Inches(0.3),
        runs=[{"text": title, "font": FONT_SANS, "size": Pt(13),
               "bold": True, "color": C_CHUMBO, "spacing": -0.3}])
    # underline under head
    add_line(s, cx + pad, cy + pad + Inches(0.4),
             cw - 2*pad, C_GRAY300, thickness=Emu(6000))
    # Items
    iy = cy + pad + Inches(0.55)
    for it in items:
        add_text(s, cx + pad, iy, cw - 2*pad, Inches(0.34),
            runs=[{"text": it, "font": FONT_SANS, "size": Pt(10.5),
                   "color": C_GRAY700}], line_spacing=1.35)
        iy += Inches(0.5)

body_text(s, M_LEFT, Inches(6.9), Inches(10), Inches(0.4),
    "O mesmo motor de IA gera ganhos simultâneos para áreas com objetivos distintos, o que facilita o caso de uso interno e a priorização da implantação.",
    size=Pt(10), bold_terms=["simultâneos"])

footer(s, 8)


# =============================================================
# SLIDE 9 · DIFERENCIAIS (8 itens 2×4)
# =============================================================
s = prs.slides.add_slide(blank_layout)
header(s, "08 · Diferenciais")
kicker(s, M_LEFT, Inches(1.05), "DIFERENCIAIS DA DOUTOR AI")
add_text(s, M_LEFT, Inches(1.35), Inches(12), Inches(1.3),
    runs=[
        {"text": "O que nos ", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "color": C_CHUMBO, "spacing": -1},
        {"text": "distingue.", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1},
    ], line_spacing=1.0)

diffs = [
    ("01", "Especialização em saúde",     "IA construída desde o início para fluxos clínicos, não adaptada de outros setores."),
    ("02", "Aplicada a fluxos reais",     "Modelagem com base em protocolos, jornadas e processos assistenciais concretos."),
    ("03", "Personalização por jornada",  "Configuração por protocolo, unidade e perfil de paciente."),
    ("04", "Integração com sistemas",     "Convive com prontuários, ERPs e ferramentas hospitalares existentes."),
    ("05", "Operação com validação humana","Humano no centro nas etapas críticas — IA como copiloto, nunca substituto."),
    ("06", "Capacidade multimodal",       "Texto, voz e dados estruturados em um mesmo motor de orquestração."),
    ("07", "Flexibilidade de adoção",     "Caminho claro de POC → piloto → expansão, com baixo atrito inicial."),
    ("08", "Segurança e governança",      "Foco em segurança da informação, rastreabilidade e conformidade regulatória."),
]
cw = (SW - M_LEFT - M_RIGHT - Inches(0.2)) / 2
ch = Inches(0.85)
cy0 = Inches(2.9)
for i, (n, t, d) in enumerate(diffs):
    row, col = divmod(i, 2)
    cx = M_LEFT + col * (cw + Inches(0.2))
    cy = cy0 + row * (ch + Inches(0.13))
    card(s, cx, cy, cw, ch)
    pad = Inches(0.17)
    # number badge
    add_rect(s, cx + pad, cy + pad, Inches(0.32), Inches(0.32),
             fill=C_BLUE_PALE, line=None)
    add_text(s, cx + pad, cy + pad, Inches(0.32), Inches(0.32),
        runs=[{"text": n, "font": FONT_MONO, "size": Pt(8.5),
               "bold": True, "color": C_BLUE_BRAND, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tx = cx + pad + Inches(0.45)
    add_text(s, tx, cy + pad - Inches(0.02), cw - pad - Inches(0.55), Inches(0.3),
        runs=[{"text": t, "font": FONT_SANS, "size": Pt(11.5),
               "bold": True, "color": C_CHUMBO, "spacing": -0.3}])
    add_text(s, tx, cy + pad + Inches(0.27), cw - pad - Inches(0.55), Inches(0.4),
        runs=[{"text": d, "font": FONT_SANS, "size": Pt(9),
               "color": C_GRAY700}], line_spacing=1.35)

footer(s, 9)


# =============================================================
# SLIDE 10 · MODELO DE IMPLANTAÇÃO (7 etapas)
# =============================================================
s = prs.slides.add_slide(blank_layout)
header(s, "09 · Implantação")
kicker(s, M_LEFT, Inches(1.05), "MODELO DE IMPLANTAÇÃO")
add_text(s, M_LEFT, Inches(1.35), Inches(12), Inches(1.3),
    runs=[
        {"text": "Como ", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "color": C_CHUMBO, "spacing": -1},
        {"text": "implementamos.", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1},
    ], line_spacing=1.0)

body_text(s, M_LEFT, Inches(2.85), Inches(11.5), Inches(0.6),
    "Adoção incremental e auditável, com início por uma frente específica de alto "
    "impacto e expansão guiada por valor comprovado.",
    size=Pt(12), bold_terms=["incremental e auditável"])

# Steps row
steps = [
    "Diagnóstico da jornada",
    "Mapeamento de fluxos e protocolos",
    "Configuração da IA",
    "Integração com sistemas",
    "POC em ambiente controlado",
    "Validação com equipe assistencial",
    "Ajustes e expansão",
]
total_w = SW - M_LEFT - M_RIGHT
sw = total_w / 7
sy = Inches(4.2)
sh = Inches(1.25)
for i, t in enumerate(steps):
    sx = M_LEFT + i * sw
    # Use 1 unified border by drawing each box; borders overlap (acceptable)
    card(s, sx, sy, sw, sh)
    pad = Inches(0.13)
    add_text(s, sx + pad, sy + pad, sw - 2*pad, Inches(0.22),
        runs=[{"text": f"{i+1:02d}", "font": FONT_MONO, "size": Pt(8),
               "color": C_BLUE_BRAND, "spacing": 1.6, "bold": True}])
    add_text(s, sx + pad, sy + pad + Inches(0.3), sw - 2*pad, sh - Inches(0.5),
        runs=[{"text": t, "font": FONT_SANS, "size": Pt(10),
               "bold": True, "color": C_CHUMBO, "spacing": -0.2}],
        line_spacing=1.2)

# Callout
cy = Inches(5.85)
ch = Inches(0.85)
card(s, M_LEFT, cy, SW - M_LEFT - M_RIGHT, ch,
     fill=C_BLUE_PALE, border=C_BLUE_PALE, left_accent=C_BLUE_BRAND)
add_text(s, M_LEFT + Inches(0.25), cy + Inches(0.13), Inches(3), Inches(0.22),
    runs=[{"text": "▸  ABORDAGEM", "font": FONT_MONO, "size": Pt(8.5),
           "color": C_BLUE_BRAND, "spacing": 1.8, "bold": True}])
add_text(s, M_LEFT + Inches(0.25), cy + Inches(0.4),
         SW - M_LEFT - M_RIGHT - Inches(0.5), Inches(0.45),
    runs=[
        {"text": "A implantação pode começar com ", "font": FONT_SANS,
         "size": Pt(11), "color": C_CHUMBO},
        {"text": "uma frente específica de alto impacto", "font": FONT_SANS,
         "size": Pt(11), "color": C_CHUMBO, "bold": True},
        {"text": ", validando valor antes da expansão para outras jornadas e unidades.",
         "font": FONT_SANS, "size": Pt(11), "color": C_CHUMBO},
    ], line_spacing=1.4)

footer(s, 10)


# =============================================================
# SLIDE 11 · VALOR ESPERADO
# =============================================================
s = prs.slides.add_slide(blank_layout)
header(s, "10 · Valor esperado")
kicker(s, M_LEFT, Inches(1.05), "VALOR ESPERADO")
add_text(s, M_LEFT, Inches(1.35), Inches(12), Inches(1.3),
    runs=[
        {"text": "O que se ", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "color": C_CHUMBO, "spacing": -1},
        {"text": "transforma.", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1},
    ], line_spacing=1.0)

# Left bullets
bullets = [
    ("Mais eficiência",        " no atendimento e em fluxos operacionais"),
    ("Redução",                " de tarefas manuais e repetitivas"),
    ("Melhor qualidade",       " da informação clínica registrada"),
    ("Padronização",           " de protocolos entre unidades e equipes"),
    ("Pertinência",            " clínica e financeira nas condutas"),
    ("Apoio",                  " consistente à tomada de decisão clínica"),
    ("Inteligência operacional"," com dados estruturados"),
    ("Escalabilidade",         " operacional com governança"),
]
bx = M_LEFT
by0 = Inches(3.05)
for i, (b, r) in enumerate(bullets):
    by = by0 + i * Inches(0.34)
    add_rect(s, bx, by + Inches(0.13), Inches(0.07), Inches(0.07),
             fill=C_BLUE_BRAND, line=None)
    add_text(s, bx + Inches(0.18), by, Inches(5.5), Inches(0.3),
        runs=[
            {"text": b, "font": FONT_SANS, "size": Pt(11.5),
             "bold": True, "color": C_CHUMBO},
            {"text": r, "font": FONT_SANS, "size": Pt(11.5), "color": C_GRAY700},
        ])

# Right chips (2x2 + 1 wide primary)
rx0 = Inches(7.5)
cw = (SW - rx0 - M_RIGHT - Inches(0.15)) / 2
ch = Inches(1.25)
chips = [
    ("01", "Eficiência",  "Menos retrabalho e mais throughput por equipe.", False),
    ("02", "Segurança",   "Padronização, rastreabilidade e governança.",      False),
    ("03", "Pertinência", "Condutas e cobranças aderentes a protocolos e contratos.", False),
    ("04", "Escala",      "Capacidade ampliada sem perda de qualidade.",      False),
]
for i, (n, t, d, prim) in enumerate(chips):
    row, col = divmod(i, 2)
    cx = rx0 + col * (cw + Inches(0.15))
    cy = Inches(3.0) + row * (ch + Inches(0.15))
    card(s, cx, cy, cw, ch, fill=C_WHITE, border=C_GRAY300)
    pad = Inches(0.18)
    add_text(s, cx + pad, cy + pad, cw - 2*pad, Inches(0.2),
        runs=[{"text": n, "font": FONT_MONO, "size": Pt(8),
               "color": C_BLUE_BRAND, "spacing": 1.6, "bold": True}])
    add_text(s, cx + pad, cy + pad + Inches(0.22), cw - 2*pad, Inches(0.4),
        runs=[{"text": t, "font": FONT_SANS, "size": Pt(17),
               "bold": True, "color": C_CHUMBO, "spacing": -0.5}],
        line_spacing=1.05)
    add_text(s, cx + pad, cy + pad + Inches(0.62), cw - 2*pad, Inches(0.5),
        runs=[{"text": d, "font": FONT_SANS, "size": Pt(9),
               "color": C_GRAY700}], line_spacing=1.4)

# Primary wide chip
py = Inches(3.0) + 2 * (ch + Inches(0.15))
pw = (SW - rx0 - M_RIGHT)
card(s, rx0, py, pw, ch, fill=C_BLUE_BRAND, border=C_BLUE_BRAND)
pad = Inches(0.22)
add_text(s, rx0 + pad, py + pad, pw - 2*pad, Inches(0.2),
    runs=[{"text": "05", "font": FONT_MONO, "size": Pt(8),
           "color": RGBColor(0xFF,0xFF,0xFF), "spacing": 1.6, "bold": True}])
add_text(s, rx0 + pad, py + pad + Inches(0.22), pw - 2*pad, Inches(0.4),
    runs=[{"text": "Inteligência operacional", "font": FONT_SANS, "size": Pt(17),
           "bold": True, "color": C_WHITE, "spacing": -0.5}],
    line_spacing=1.05)
add_text(s, rx0 + pad, py + pad + Inches(0.62), pw - 2*pad, Inches(0.5),
    runs=[{"text": "Dados estruturados para decisão clínica e de gestão.",
           "font": FONT_SANS, "size": Pt(9.5),
           "color": C_WHITE}], line_spacing=1.4)

# Bottom disclaimer
add_text(s, M_LEFT, SH - Inches(0.7),
         SW - M_LEFT - M_RIGHT, Inches(0.25),
    runs=[{"text": "Benefícios esperados, sustentados pelo desenho da plataforma e pelas jornadas implantadas. Métricas absolutas dependem do escopo e da maturidade de cada instituição.",
           "font": FONT_SANS, "size": Pt(8.5), "color": C_GRAY500}],
    line_spacing=1.4)

footer(s, 11)


# =============================================================
# SLIDE 12 · FECHAMENTO & PRÓXIMOS PASSOS
# =============================================================
s = prs.slides.add_slide(blank_layout)
header(s, "11 · Próximos passos")
kicker(s, M_LEFT, Inches(1.05), "FECHAMENTO & PRÓXIMOS PASSOS")
add_text(s, M_LEFT, Inches(1.35), Inches(12), Inches(1.3),
    runs=[
        {"text": "Como ", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "color": C_CHUMBO, "spacing": -1},
        {"text": "avançar.", "font": FONT_SANS, "size": Pt(36),
         "bold": True, "italic": True, "color": C_BLUE_BRAND, "spacing": -1},
    ], line_spacing=1.0)

steps_next = [
    ("01 · ESCOPO",         "Definir a jornada prioritária"),
    ("02 · FRENTE INICIAL", "Escolher unidade, canal ou fluxo de partida"),
    ("03 · INSUMOS",        "Levantar protocolos e scripts existentes"),
    ("04 · TÉCNICO",        "Alinhar integrações necessárias"),
    ("05 · MÉTRICAS",       "Definir critérios de sucesso"),
    ("06 · EXECUÇÃO",       "Iniciar pela frente escolhida — POC ou implantação direta"),
]
cw = (SW - M_LEFT - M_RIGHT - Inches(0.3)) / 3
ch = Inches(0.85)
cy0 = Inches(2.9)
for i, (n, t) in enumerate(steps_next):
    row, col = divmod(i, 3)
    cx = M_LEFT + col * (cw + Inches(0.15))
    cy = cy0 + row * (ch + Inches(0.15))
    card(s, cx, cy, cw, ch)
    pad = Inches(0.18)
    add_text(s, cx + pad, cy + pad, cw - 2*pad, Inches(0.22),
        runs=[{"text": n, "font": FONT_MONO, "size": Pt(8.5),
               "color": C_BLUE_BRAND, "spacing": 1.6, "bold": True}])
    add_text(s, cx + pad, cy + pad + Inches(0.28), cw - 2*pad, Inches(0.4),
        runs=[{"text": t, "font": FONT_SANS, "size": Pt(11.5),
               "bold": True, "color": C_CHUMBO, "spacing": -0.3}],
        line_spacing=1.15)

# Final dark message
fy = Inches(5.05)
fh = Inches(1.55)
add_rect(s, M_LEFT, fy, SW - M_LEFT - M_RIGHT, fh, fill=C_CHUMBO, line=None)
add_text(s, M_LEFT + Inches(0.35), fy + Inches(0.2),
         SW - M_LEFT - M_RIGHT - Inches(0.7), Inches(0.22),
    runs=[{"text": "▸  MENSAGEM FINAL", "font": FONT_MONO, "size": Pt(8.5),
           "color": C_BLUE_LIGHT, "spacing": 1.8, "bold": True}])
add_text(s, M_LEFT + Inches(0.35), fy + Inches(0.5),
         SW - M_LEFT - M_RIGHT - Inches(0.7), Inches(1.0),
    runs=[
        {"text": "A Doutor AI ajuda instituições de saúde a transformar jornadas complexas em fluxos ",
         "font": FONT_SANS, "size": Pt(16), "bold": True, "color": C_WHITE,
         "spacing": -0.4},
        {"text": "mais inteligentes, eficientes e centrados no paciente.",
         "font": FONT_SANS, "size": Pt(16), "bold": True, "italic": True,
         "color": C_BLUE_LIGHT, "spacing": -0.4},
    ], line_spacing=1.3)

footer(s, 12)


# =============================================================
# SAVE
# =============================================================
out_path = "/home/user/-/doutor_ai_apresentacao_institucional.pptx"
prs.save(out_path)
print(f"OK -> {out_path}")
