#!/usr/bin/env python3
"""Gera a planilha de checklist comercial Doutor-AI."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
wb.remove(wb.active)


# ── helpers ──────────────────────────────────────────────────────────────────

def fill(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

def border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def bold_border():
    s = Side(style="medium", color="888888")
    return Border(left=s, right=s, top=s, bottom=s)


# ── data ─────────────────────────────────────────────────────────────────────

PHASES = [
    {
        "name": "AQUISIÇÃO",
        "icon": "🔍",
        "header_color": "1F4E79",
        "macro_color": "2E75B6",
        "light_color": "BDD7EE",
        "row_alt": "EBF3FB",
        "macros": [
            {
                "name": "Prospecção & Qualificação",
                "activities": [
                    ("Determinar metas e métricas comerciais", "AE", "D+0 → D+2"),
                    ("Pesquisar conta-alvo e mapear stakeholders (CEO, COO, CIO)", "SDR", "D+0 → D+2"),
                    ("Executar cadência outbound (e-mail + LinkedIn + telefone)", "SDR", "D+0 → D+5"),
                    ("Agendar discovery call com CMO/sponsor", "SDR · AE", "D+5 → D+7"),
                    ("Registrar account plan e próximos passos no CRM", "SDR", "D+7"),
                    ("Revisão de CRM e forecast preditivo de win rate (modelo ML)", "AE", "Backlog"),
                ],
            },
            {
                "name": "Pré-venda Clínica & Demo",
                "activities": [
                    ("Coletar brief técnico-clínico do prospect (durante call)", "BDR", "D+0"),
                    ("Customizar deck clínico com cases do segmento", "BDR", "D+2 → D+4"),
                    ("Conduzir demo guiada (AE + Clinical Specialist)", "BDR", "D+5 → D+7"),
                    ("Responder DC clínico-regulatório", "BDR", "≤ 5 dias úteis"),
                    ("Obter sign-off do CMO Doutor-AI no whitepaper aplicado", "Medical Director", "D+10"),
                    ("Construir relatório de cliente 300d com dados anonimizados", "Medical Director", "D+15 → D+45"),
                    ("Painel interativo de simulação de impacto financeiro", "Medical Director", "Backlog"),
                    ("Montar proposta com módulos, volumetria e pricing", "AE", "D+0 → D+5"),
                ],
            },
        ],
        "handoff_label": "✅  HANDOFF  →  Contratação & Fechamento",
        "handoff_criteria": "Proposta montada e validada  ·  Brief clínico coletado  ·  Sign-off CMO obtido",
    },
    {
        "name": "CONTRATAÇÃO & FECHAMENTO",
        "icon": "📝",
        "header_color": "843C0C",
        "macro_color": "C55A11",
        "light_color": "FCE4D6",
        "row_alt": "FFF0E8",
        "macros": [
            {
                "name": "Negociação & Assinatura",
                "activities": [
                    ("Apresentar e enviar a proposta ao cliente", "AE", "D+5 → D+8"),
                    ("Negociar contrato com cliente", "AE", "D+8 → D+15"),
                    ("Assinar Contrato + Order Form e lançar booking no CRM/ERP", "AE", "D+20"),
                ],
            },
        ],
        "handoff_label": "✅  HANDOFF  →  Onboarding",
        "handoff_criteria": "Contrato assinado  ·  Order Form assinado  ·  Booking registrado no CRM/ERP",
    },
    {
        "name": "ONBOARDING",
        "icon": "🚀",
        "header_color": "375623",
        "macro_color": "538135",
        "light_color": "E2EFDA",
        "row_alt": "F4FAEE",
        "macros": [
            {
                "name": "Pós-fechamento & Handover Interno",
                "activities": [
                    ("AE cria linha na Planilha de Gestão de Contratos (1 linha/cliente)", "AE · CSM", "D+0 → D+1"),
                    ("CS cria linha na Planilha de Planejamento de Recebimentos", "AE · CSM", "D+1"),
                    ("AE envia e-mail de boas-vindas ao cliente (sponsor + champions)", "AE · CSM", "D+1"),
                    ("Enviar e-mail interno de handover (AE → CS · Impl · Med · Ops · Fin)", "AE · CSM", "D+1"),
                    ("Head Implantação preenche matriz de capacity atual", "AE · CSM", "D+1 → D+2"),
                    ("Agendar reunião interna de entrega — área de alocação de squads", "AE · CSM", "D+2 → D+3"),
                    ("Abrir protocolo do cliente (pasta Drive + canal Slack #cli-[nome])", "AE · CSM", "D+2 → D+3"),
                    ("Preparar questionário pré-Kickoff (Clínica / Técnica / Produto / CS)", "AE · CSM", "D+3"),
                    ("Enviar questionário pré-kickoff ao cliente (sponsor + champions)", "AE · CSM", "D+3 → D+4"),
                    ("Cliente responde + squad consolida e revisa respostas por área", "AE · CSM", "D+4 → D+6"),
                    ("Agendar Kickoff oficial com sponsor + diretoria do cliente (4h, presencial)", "AE · CSM", "D+6 → D+8"),
                ],
            },
        ],
        "handoff_label": "✅  HANDOFF  →  CS + Implantação",
        "handoff_criteria": "Kickoff agendado  ·  Questionário respondido e revisado  ·  Squad alocado  ·  Canal Slack aberto",
    },
]


# ── Sheet 1 : Pipeline Geral ──────────────────────────────────────────────────

ws_pip = wb.create_sheet("🏠 Pipeline Geral")

# title
ws_pip.merge_cells("A1:M1")
c = ws_pip["A1"]
c.value = "PIPELINE COMERCIAL — ACOMPANHAMENTO DE CLIENTES"
c.font = Font(name="Calibri", bold=True, size=15, color="FFFFFF")
c.fill = fill("1F4E79")
c.alignment = Alignment(horizontal="center", vertical="center")
ws_pip.row_dimensions[1].height = 34

ws_pip.merge_cells("A2:M2")
c = ws_pip["A2"]
c.value = (
    "Acompanhe o estágio de cada cliente e identifique quando realizar o handoff para a próxima área."
)
c.font = Font(name="Calibri", italic=True, size=9, color="595959")
c.fill = fill("EBF3FB")
c.alignment = Alignment(horizontal="center", vertical="center")
ws_pip.row_dimensions[2].height = 18

headers = [
    "Cliente / Hospital",
    "Operadora / Segmento",
    "AE Responsável",
    "SDR",
    "Data Entrada",
    "Fase Atual",
    "% Aquisição",
    "% Contratação",
    "% Onboarding",
    "Bloqueio?",
    "Motivo do Bloqueio",
    "Data Handoff",
    "Observações",
]
ws_pip.row_dimensions[3].height = 36
for ci, h in enumerate(headers, 1):
    c = ws_pip.cell(row=3, column=ci, value=h)
    c.font = Font(name="Calibri", bold=True, size=9, color="FFFFFF")
    c.fill = fill("1F4E79")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border()

col_widths = [28, 22, 18, 14, 13, 24, 13, 14, 13, 11, 28, 13, 32]
for i, w in enumerate(col_widths, 1):
    ws_pip.column_dimensions[get_column_letter(i)].width = w

fase_dv = DataValidation(
    type="list",
    formula1='"Aquisição,Contratação & Fechamento,Onboarding,CS,Implantação"',
    allow_blank=True,
)
ws_pip.add_data_validation(fase_dv)
bloq_dv = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
ws_pip.add_data_validation(bloq_dv)

for r in range(4, 54):
    bg = "FAFAFA" if r % 2 == 0 else "FFFFFF"
    ws_pip.row_dimensions[r].height = 22
    for ci in range(1, 14):
        c = ws_pip.cell(row=r, column=ci)
        c.fill = fill(bg)
        c.border = border()
        c.font = Font(name="Calibri", size=9)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    fase_dv.add(ws_pip.cell(row=r, column=6))
    bloq_dv.add(ws_pip.cell(row=r, column=10))

ws_pip.freeze_panes = "A4"


# ── Sheet 2 : Checklist (template por cliente) ────────────────────────────────

ws_cl = wb.create_sheet("📋 Checklist - Cliente")

ws_cl.column_dimensions["A"].width = 20   # status
ws_cl.column_dimensions["B"].width = 52   # atividade
ws_cl.column_dimensions["C"].width = 20   # responsável
ws_cl.column_dimensions["D"].width = 16   # prazo previsto
ws_cl.column_dimensions["E"].width = 16   # data conclusão
ws_cl.column_dimensions["F"].width = 34   # observações

# title
ws_cl.merge_cells("A1:F1")
c = ws_cl["A1"]
c.value = "CHECKLIST COMERCIAL — JORNADA DO CLIENTE"
c.font = Font(name="Calibri", bold=True, size=15, color="FFFFFF")
c.fill = fill("1F4E79")
c.alignment = Alignment(horizontal="center", vertical="center")
ws_cl.row_dimensions[1].height = 34

# client info row 2
info = [
    ("A2", "B2", "Cliente:"),
    ("C2", "D2", "AE Responsável:"),
    ("E2", "F2", "Data de Início:"),
]
ws_cl.row_dimensions[2].height = 22
ws_cl.row_dimensions[3].height = 22

label_font = Font(name="Calibri", bold=True, size=9, color="FFFFFF")
label_fill = fill("2E75B6")
input_fill = fill("EBF3FB")
label_align = Alignment(horizontal="left", vertical="center", indent=1)

for lc, ic, txt in info:
    c = ws_cl[lc]
    c.value = txt; c.font = label_font; c.fill = label_fill; c.alignment = label_align
    c = ws_cl[ic]
    c.fill = input_fill; c.border = border()

info3 = [
    ("A3", "B3", "Hospital/Operadora:"),
    ("C3", "D3", "SDR:"),
    ("E3", "F3", "Segmento ICP:"),
]
for lc, ic, txt in info3:
    c = ws_cl[lc]
    c.value = txt; c.font = label_font; c.fill = label_fill; c.alignment = label_align
    c = ws_cl[ic]
    c.fill = input_fill; c.border = border()

# status dropdown
status_dv = DataValidation(
    type="list",
    formula1='"⬜ Pendente,🔄 Em andamento,✅ Concluído,⏭ N/A"',
    allow_blank=True,
)
ws_cl.add_data_validation(status_dv)

row = 5  # start writing checklist content

def phase_header(ws, r, phase):
    ws.row_dimensions[r].height = 30
    ws.merge_cells(f"A{r}:F{r}")
    c = ws.cell(row=r, column=1, value=f"  {phase['icon']}  {phase['name']}")
    c.font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    c.fill = fill(phase["header_color"])
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    return r + 1

def macro_header(ws, r, name, color):
    ws.row_dimensions[r].height = 22
    ws.merge_cells(f"A{r}:F{r}")
    c = ws.cell(row=r, column=1, value=f"   {name}")
    c.font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
    c.fill = fill(color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=2)
    return r + 1

def col_headers(ws, r, light):
    ws.row_dimensions[r].height = 20
    labels = ["Status", "Atividade", "Responsável", "Prazo Previsto", "Data Conclusão", "Observações"]
    for ci, lbl in enumerate(labels, 1):
        c = ws.cell(row=r, column=ci, value=lbl)
        c.font = Font(name="Calibri", bold=True, size=8, color="404040")
        c.fill = fill(light)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border()
    return r + 1

def activity_row(ws, r, act, actor, prazo, row_alt, status_dv):
    ws.row_dimensions[r].height = 20
    bg = row_alt if r % 2 == 0 else "FFFFFF"

    vals = ["⬜ Pendente", act, actor, prazo, "", ""]
    aligns = ["center", "left", "center", "center", "center", "left"]
    editable = [False, False, False, False, True, True]

    for ci, (val, aln, ed) in enumerate(zip(vals, aligns, editable), 1):
        c = ws.cell(row=r, column=ci, value=val)
        c.font = Font(name="Calibri", size=9)
        c.fill = fill("FFFCE5") if ed else fill(bg)
        c.alignment = Alignment(horizontal=aln, vertical="center", indent=(1 if aln == "left" else 0))
        c.border = border()

    status_dv.add(ws.cell(row=r, column=1))
    return r + 1

def handoff_row(ws, r, label, criteria):
    # blank gap
    ws.row_dimensions[r].height = 8
    r += 1

    ws.row_dimensions[r].height = 28
    ws.merge_cells(f"A{r}:F{r}")
    c = ws.cell(row=r, column=1, value=label)
    c.font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    c.fill = fill("00884A")
    c.alignment = Alignment(horizontal="center", vertical="center")
    r += 1

    ws.row_dimensions[r].height = 20
    ws.merge_cells(f"A{r}:F{r}")
    c = ws.cell(row=r, column=1, value=f"Critérios de conclusão:  {criteria}")
    c.font = Font(name="Calibri", italic=True, size=8, color="1A4314")
    c.fill = fill("D5EDCA")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=2)
    c.border = border()
    r += 1

    ws.row_dimensions[r].height = 10
    return r + 1

for phase in PHASES:
    row = phase_header(ws_cl, row, phase)
    for macro in phase["macros"]:
        row = macro_header(ws_cl, row, macro["name"], phase["macro_color"])
        row = col_headers(ws_cl, row, phase["light_color"])
        for act, actor, prazo in macro["activities"]:
            row = activity_row(ws_cl, row, act, actor, prazo, phase["row_alt"], status_dv)
    row = handoff_row(ws_cl, row, phase["handoff_label"], phase["handoff_criteria"])

ws_cl.freeze_panes = "A5"


# ── Sheet 3 : Instruções ──────────────────────────────────────────────────────

ws_i = wb.create_sheet("ℹ️ Instruções")

ws_i.column_dimensions["A"].width = 16
ws_i.column_dimensions["B"].width = 60
ws_i.column_dimensions["C"].width = 5

ws_i.merge_cells("A1:B1")
c = ws_i["A1"]
c.value = "COMO USAR ESTE CHECKLIST"
c.font = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
c.fill = fill("1F4E79")
c.alignment = Alignment(horizontal="center", vertical="center")
ws_i.row_dimensions[1].height = 30

steps = [
    ("PASSO 1", "No Pipeline Geral, adicione cada cliente novo como uma linha."),
    ("PASSO 2", "Duplique a aba '📋 Checklist - [Cliente]' e renomeie com o nome do cliente."),
    ("PASSO 3", "Preencha o cabeçalho da aba: nome do cliente, AE, SDR, data de início."),
    ("PASSO 4", "À medida que o cliente avança, atualize o Status de cada atividade via dropdown."),
    ("PASSO 5", "Registre a Data Conclusão e Observações nas colunas amarelas."),
    ("PASSO 6", "Quando todos os itens de uma fase estiverem ✅ Concluídos, execute o handoff."),
    ("PASSO 7", "Atualize a coluna 'Fase Atual' e 'Data Handoff' no Pipeline Geral."),
]

legend = [
    ("⬜ Pendente", "Atividade ainda não iniciada"),
    ("🔄 Em andamento", "Atividade em execução no momento"),
    ("✅ Concluído", "Atividade finalizada — preencha a Data Conclusão"),
    ("⏭ N/A", "Não aplicável para este cliente/contexto"),
]

r = 3

def section_title(ws, r, txt, color):
    ws.merge_cells(f"A{r}:B{r}")
    c = ws.cell(row=r, column=1, value=txt)
    c.font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
    c.fill = fill(color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r].height = 22
    return r + 1

r = section_title(ws_i, r, "PASSO A PASSO", "2E75B6")
for step, desc in steps:
    ws_i.row_dimensions[r].height = 20
    c = ws_i.cell(row=r, column=1, value=step)
    c.font = Font(name="Calibri", bold=True, size=9, color="1F4E79")
    c.fill = fill("BDD7EE")
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border()
    c = ws_i.cell(row=r, column=2, value=desc)
    c.font = Font(name="Calibri", size=9)
    c.fill = fill("EBF3FB")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c.border = border()
    r += 1

r += 1
r = section_title(ws_i, r, "LEGENDA DE STATUS", "375623")
for status, desc in legend:
    ws_i.row_dimensions[r].height = 20
    c = ws_i.cell(row=r, column=1, value=status)
    c.font = Font(name="Calibri", bold=True, size=9)
    c.fill = fill("E2EFDA")
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border()
    c = ws_i.cell(row=r, column=2, value=desc)
    c.font = Font(name="Calibri", size=9)
    c.fill = fill("F4FAEE")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c.border = border()
    r += 1

r += 1
r = section_title(ws_i, r, "FASES & HANDOFFS", "843C0C")
handoff_info = [
    ("Aquisição → Contratação", "Proposta montada · Brief clínico coletado · Sign-off CMO obtido"),
    ("Contratação → Onboarding", "Contrato assinado · Order Form assinado · Booking no CRM/ERP"),
    ("Onboarding → CS + Implantação", "Kickoff agendado · Questionário respondido · Squad alocado · Slack aberto"),
]
for fase, crit in handoff_info:
    ws_i.row_dimensions[r].height = 20
    c = ws_i.cell(row=r, column=1, value=fase)
    c.font = Font(name="Calibri", bold=True, size=9, color="FFFFFF")
    c.fill = fill("C55A11")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border()
    c = ws_i.cell(row=r, column=2, value=crit)
    c.font = Font(name="Calibri", size=9)
    c.fill = fill("FFF0E8")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    c.border = border()
    ws_i.row_dimensions[r].height = 24
    r += 1


# ── Save ──────────────────────────────────────────────────────────────────────

path = "/home/user/-/checklist_comercial_doutor_ai.xlsx"
wb.save(path)
print(f"Arquivo salvo: {path}")
